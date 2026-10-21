"""followup.py — Daftar follow-up baris anomali (v3.2.0) — ala repo Waste.

Konsep: user mencentang baris mencurigakan (kolom ✓) di tabel DEVIASI GR /
DEVIASI PCS / LAPORAN DEVIASI → baris terkumpul di tab "☑ Follow Up" sebagai
SNAPSHOT baris (nilai apa adanya saat dicentang). Daftar menetap — TIDAK ikut
filter/H aktif; tertutup aplikasi pun aman (followup.json).

v3.9.6 (permintaan user: "format Follow UP seperti file sample"):
- Format tabel + export mengikuti sheet Follow UP workbook asli user:
  Resto · Nama Barang · SO per tanggal (H-2..H+2 — heading angka tanggal,
  mis. H=25 → 23/24/25/26/27 persis sample "22..27") · Sistem · Deviasi ·
  Loss Surplus to BOM · Keterangan · kolom kerja FU: SO · DEVIASI ·
  Selisih · Selisih Deviasi · Cek FU · PIC (+ pelacakan app: Asal ·
  Status TL · Ditambahkan).
- SNAPSHOT_FIELDS diperluas: so_m2/so_m1/so_p1/so_p2 (SO per tanggal) +
  loss_bom (Loss Surplus to BOM = Dev to BOM% tab Deviasi). Snapshot
  LAMA (tanpa field baru) tetap dimuat — kolom terkait tampil kosong.

Struktur (pola waste_anomaly/followup.py):
- ``row_fingerprint(fields)`` — md5 identitas baris: Asal|Resto|Barang|Satuan|
  H|SO|Sistem|Deviasi|CEK. Stabil lintas sesi (str(float) dinormalisasi).
- ``FollowUpStore`` — dict {fp: snapshot} + status tindak lanjut (resolved):
  - toggle(fp, row) → tambah/hapus, return status baru
  - save/load (atomic write + backup .bad utk file korup; pulihan file
    korup kini TERLIHAT — flag ``last_recovery`` + [WARNING] log + dialog
    bila GUI hidup, FIX 2f-07)
  - gagal persist → flag ``last_save_error`` utk view (FIX 2f-08)
  - export_xlsx(path) → sheet "FOLLOW UP" format file sample + kolom
    kerja FU (diisi petugas; nilai tersimpan ikut terbawa); tulis ATOMIK
    tmp+os.replace (FIX 2f-04 — file export lama tak tertimpa parsial)
  (v3.8: import_workbook DIHAPUS — fitur "Impor Hasil TL" tidak perlu
  lagi, permintaan user; resolved lama tetap dimuat agar status/PIC
  yang sudah tersimpan sebelumnya tetap tampil.)
- Modul murni — tanpa Tk di IMPORT-TIME (dipakai view + app + test
  headless); dialog pemulihan 2f-07 meng-import tkinter LAZY dan hanya
  tampil bila root Tk default hidup (mode GUI).
"""
from __future__ import annotations

import hashlib
import json
import logging
import os
from datetime import datetime
from pathlib import Path

import pandas as pd

from ..config import app_base_dir as _app_base_dir  # v3.9.62 — frozen-aware

_LOG = logging.getLogger("stock_monitoring.followup")

# Snapshot row keys (persis kolom tabel Follow Up, tanpa ✓/status TL).
# v3.9.6 (permintaan user: "format Follow UP seperti file sample" — sheet
# Follow UP workbook asli user): snapshot diperluas — SO per tanggal
# (H-2..H+2) + Loss Surplus to BOM (= Dev to BOM% tab Deviasi).
SNAPSHOT_FIELDS = [
    "asal", "resto", "barang", "satuan", "aktif", "h",
    "so", "sistem", "deviasi", "cek",
    "so_m2", "so_m1", "so_p1", "so_p2", "loss_bom",
    # v3.9.54 (audit 8-a P1-3) — nama versi APP baris asal KOMPARASI
    # ('resto' snap = bentuk pendek ternormalisasi, 'barang' = nama FILE
    # → popup detail exact-match data app = 0 baris / popup KOSONG).
    # Snap lama tanpa field ini → fallback 'barang'/'resto' (perilaku
    # lama). TIDAK mengubah fingerprint (row_fingerprint md5 9 field).
    "barang_app", "resto_app",
]
# Kolom tanggal SO di tabel/export (posisi H relatif — heading dinamis
# menampilkan angka tanggal, mis. H=25 → 23/24/25/26/27)
DATE_COL_IDS = ["T-2", "T-1", "T", "T+1", "T+2"]

# Kolom serah-terima di export (diisi petugas follow-up) — mengikuti
# sheet Follow UP file sample: SO · DEVIASI · Selisih · Selisih
# Deviasi · Cek FU (+ PIC utk pelacakan app).
HANDOVER_COLS = ["SO", "DEVIASI", "Selisih", "Selisih Deviasi",
                  "Cek FU", "PIC"]

# Nama kolom sheet FOLLOW UP (export) — urutan PERSIS sheet Follow UP
# file sample user: Resto · Nama Barang · [SO tgl H-2..H+2] · Sistem ·
# Deviasi · Loss Surplus to BOM · Keterangan · [kerja FU] (+ kolom
# pelacakan app: Asal · Status TL · Ditambahkan).
EXPORT_COLS = [
    "Resto", "Nama Barang",
] + DATE_COL_IDS + [
    "Sistem", "Deviasi", "Loss Surplus to BOM", "Keterangan",
] + HANDOVER_COLS + [
    "Asal", "Status TL", "Ditambahkan",
]

# Pergeseran kolom tanggal relatif H (T-2 → h−2 … T+2 → h+2)
DATE_COL_OFFSETS = dict(zip(DATE_COL_IDS, (-2, -1, 0, 1, 2)))

CHECK_ON, CHECK_OFF = "☑", "☐"

# FIX 6d1-04 (audit 6-d1 LOW) — 7 label anomali CEK KANONIK (mirror
# deviasi_view._CEK_CHECK_LABELS — READ ONLY, jangan uba dua-duanya
# sekaligus tanpa koordinasi): snapshot 'cek' baris asal DEVIASI dapat
# berisi segmen ANALISA lalu→kini ('Berubah Arah' / 'Peningkatan
# Tertinggi +p%' / 'Penurunan Tertinggi -p%') bila baris TIDAK ber-label
# anomali — footer CEK tab Deviasi tidak menghitungnya sbg anomali.
# Display Follow Up menandainya 'Analisa: …' supaya dua tab konsisten
# (bukan tertukar dgn label anomali).
_CEK_CANONICAL_LABELS = (
    "Promix 04", "RI Balancing", "Kosong/Non Aktif", "Beda Tanggal",
    "Hitung Manual", "Turun Drastis", "ACR to BOM",
)
_ANALISA_PREFIXES = ("Berubah Arah", "Peningkatan Tertinggi",
                     "Penurunan Tertinggi")


def _cek_display_text(v) -> str:
    """FIX 6d1-04 — teks kolom Keterangan (CEK) Follow Up: label
    kanonik tampil apa adanya; segmen ANALISA diberi prefiks
    'Analisa: ' (baris tsb BUKAN anomali CEK — definisi footer/popup
    tab Deviasi); label lain (status LAPORAN-COM: MINUS / > MEDIAN /
    MINUS + > MEDIAN) tetap apa adanya — konsisten dgn tab asalnya."""
    t = str(v or "").strip()
    if not t:
        return "-"
    if t in _CEK_CANONICAL_LABELS:
        return t
    if t.startswith(_ANALISA_PREFIXES):
        return f"Analisa: {t}"
    return t

# Path (dynamis — pola config.json di app.py)
# v3.9.62 — frozen-aware (app_base_dir): mode PyInstaller → folder .exe.
_BASE = _app_base_dir()


def followup_path() -> Path:
    return _BASE / "followup.json"


def resolved_path() -> Path:
    return _BASE / "followup_resolved.json"


def _norm_num(v) -> str:
    """Normalisasi angka utk fingerprint: 5 == 5.0 == '5.0' → '5'."""
    try:
        f = float(v)
        if f == int(f):
            return str(int(f))
        return repr(round(f, 6))
    except (TypeError, ValueError):
        return "" if v is None else str(v).strip()


def _pyval(v):
    """Konversi nilai snapshot → tipe Python murni (numpy int64/float64
    → int/float; pd.NA → None) supaya JSON-serializable."""
    if v is None:
        return ""
    if v is pd.NA:
        return ""
    item = getattr(v, "item", None)   # numpy scalar / pd.Timestamp
    if callable(item):
        try:
            return v.item()
        except Exception:
            pass
    if isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def row_fingerprint(row: dict) -> str:
    """md5 identitas baris (stabil lintas sesi).

    v3.9.69 — hardening (audit eksternal ke-2, risiko rendah): tiap part
    di-ESCAPE dulu (``\\`` → ``\\\\``, ``|`` → ``\\|``) sebelum di-join —
    menutup bentrok teoretis "A|B"+"C" ≡ "A"+"B|C" (nilai memuat pipe).
    BACKWARD-COMPAT: utk data bebas-pipe (semua data nyata saat ini —
    audit: "belum terbukti di data sekarang") hash TIDAK berubah sama
    sekali, jadi followup.json lama tetap cocok; hash berubah HANYA utk
    nilai yang memang ber-pipe (yang dulu ambigu sejak awal)."""
    def _esc(v: str) -> str:
        return v.replace("\\", "\\\\").replace("|", "\\|")

    parts = [
        _esc(str(row.get("asal", "") or "").strip()),
        _esc(str(row.get("resto", "") or "").strip()),
        _esc(str(row.get("barang", "") or "").strip()),
        _esc(str(row.get("satuan", "") or "").strip()),
        _esc(_norm_num(row.get("h"))),
        _esc(_norm_num(row.get("so"))),
        _esc(_norm_num(row.get("sistem"))),
        _esc(_norm_num(row.get("deviasi"))),
        _esc(str(row.get("cek", "") or "").strip()),
    ]
    key = "|".join(parts)
    return hashlib.md5(key.encode("utf-8", "surrogatepass")).hexdigest()


def _backup_corrupt(p: Path) -> None:
    """File korup di-rename .bad (tidak hilang) — pola Waste."""
    try:
        os.replace(p, p.with_suffix(p.suffix + ".bad"))
    except OSError:
        pass


def _notify_recovery(store, p: Path, reason: str) -> None:
    """v3.9.59 — FIX 2f-07 (audit 2-f MED): pulihan file korup kini
    TERLIHAT oleh user (dulu SENYAP: file di-rename .bad + daftar mulai
    kosong tanpa pesan — _LOG.warning jatuh ke stderr yang dibuang
    pythonw/MULAI.bat, user kehilangan seluruh daftar follow-up tanpa
    tahu cadangan .bad ada).

    Tiga kanal (aman utk jalur headless/test — tanpa root Tk dialog
    TIDAK pernah dicoba):
    1. ``store.last_recovery`` = pesan ringkas — view/app (listener
       v3.9.42) bisa men-surface flash status (API utk view).
    2. print [WARNING] — stdout app dialihkan ke log internal tab 🚀
       Proses (pola save-gagal followup v3.9.42).
    3. messagebox showwarning HANYA bila root Tk default hidup (mode
       GUI; dibungkus try — store bisa dibuat sebelum mainloop/di
       thread non-main)."""
    bad = str(p) + ".bad"
    msg = (f"File {p.name} rusak ({reason}). Cadangan: {bad}. "
           "Daftar dimulai kosong.")
    if store is not None:
        store.last_recovery = msg
    print(f"[WARNING] followup: {msg}")
    try:
        import tkinter as tk
    except ImportError:
        return  # Tk tak tersedia (jalur headless) — cukup flag + log
    if getattr(tk, "_default_root", None) is None:
        return  # tanpa root default (test headless) — JANGAN dialog
    try:
        from tkinter import messagebox
        messagebox.showwarning("Follow-Up Rusak", msg)
    except (tk.TclError, RuntimeError):
        _LOG.debug("dialog followup-rusak gagal (non-main thread?)",
                   exc_info=True)


def _fmt_added(ts) -> str:
    """v3.9.59 — FIX 2f-15 (audit 2-f): kolom 'Ditambahkan' export pakai
    kanon timestamp pendek app '%d-%m-%Y %H:%M' (dulu ISO
    '2025-11-12 08:30' — gaya ke-4 yang tak konsisten dgn permukaan
    lain). Snapshot lama/non-ISO → fallback tampilan lama."""
    s = str(ts or "").strip()
    if not s:
        return ""
    try:
        return datetime.strptime(s[:19], "%Y-%m-%dT%H:%M:%S").strftime(
            "%d-%m-%Y %H:%M")
    except ValueError:
        return s[:16].replace("T", " ")


def _atomic_save(path: Path, data: dict) -> bool:
    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        os.replace(tmp, path)
        return True
    except OSError as e:
        _LOG.warning("gagal simpan %s: %s", path.name, e)
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass
        return False


class FollowUpStore:
    """Store follow-up: snapshot baris tercentang + status tindak lanjut."""

    def __init__(self):
        self.rows: dict[str, dict] = {}      # fp → snapshot
        self.resolved: dict[str, dict] = {}  # fp → {PIC, KATEGORI, TINDAKAN, ts}
        self._listeners: list = []
        # v3.9.59 — FIX 2f-07: pesan pulihan file korup TERAKHIR (None =
        # tak ada) — dibaca view/app utk flash status (lihat
        # _notify_recovery).
        self.last_recovery: str | None = None
        # v3.9.59 — FIX 2f-08: pesan gagal persist TERAKHIR (None = sehat)
        # — dibaca view/listener setelah toggle/clear (lihat toggle).
        self.last_save_error: str | None = None
        self.load()

    # -- sinkronisasi UI ------------------------------------------------------
    def add_listener(self, cb) -> None:
        """cb(store) dipanggil setiap perubahan (toggle/clear)."""
        if cb not in self._listeners:
            self._listeners.append(cb)

    def remove_listener(self, fn) -> None:
        """v3.9.42 — audit 61-c/61-d (F5 P3): lepas listener (defensif,
        idempoten — fn yang tak terdaftar diabaikan senyap). Store hanya
        punya add_listener: view yang di-destroy / direbuild tidak bisa
        berhenti berlangganan → callback _notify mendarat di widget mati
        (traceback noise) + listener menumpuk seumur app. API disiapkan
        utk dipakai view pemanggil add_listener saat teardown (dipanggil
        di deviasi_view / komparasi_deviasi_view — dikerjakan pusat).

        (v3.9.33 48-e FIX 11 memakai guard winfo_exists di sisi view —
        remove_listener memberi akar fix, keduanya kompatibel.)"""
        try:
            self._listeners.remove(fn)
        except ValueError:
            pass  # fn memang tak terdaftar (defensif)

    def _notify(self):
        for cb in list(self._listeners):
            try:
                cb(self)
            except Exception:
                pass

    # -- query ------------------------------------------------------------------
    def has(self, fp: str) -> bool:
        return fp in self.rows

    def __len__(self) -> int:
        return len(self.rows)

    def resolved_count(self) -> int:
        return sum(1 for fp in self.rows if fp in self.resolved)

    # -- mutasi -------------------------------------------------------------------
    def toggle(self, fp: str, row: dict) -> bool:
        """Tambah/hapus baris. Return True bila kini TERDAFTAR.

        v3.9.42 — audit 61-c/61-d (F5 P3): hasil save()/save_resolved()
        dulu DIABAIKAN — persist gagal (file terkunci / disk penuh)
        senyap: user yakin baris ter-uncheck padahal restart
        mengembalikannya. Kini gagal → print [WARNING] (stdout app
        dialihkan ke log internal; JANGAN messagebox — context bisa
        bukan main thread). Perubahan tetap di memori (degrade soft).
        v3.9.59 — FIX 2f-08: gagal juga men-set flag
        ``self.last_save_error`` SEBELUM _notify() supaya view yang
        mendengarkan store bisa men-surface flash status di tempat
        user mencentang (log tab Proses saja terlalu jauh)."""
        if fp in self.rows:
            del self.rows[fp]
            # v3.9.33 — AUDIT FIX (48-e FIX 9): resolved dihapus dari
            # memori tapi TIDAK ditulis ke followup_resolved.json → fp
            # orphan tetap hidup di file dan muncul lagi setelah restart.
            # Simpan sekarang (baris yang ditambah ulang mulai bersih).
            self.resolved.pop(fp, None)
            ok_resolved = self.save_resolved()
            ok = self.save()
            self.last_save_error = None if (ok and ok_resolved) else (
                "gagal menyimpan status followup (file terkunci/disk?) "
                "— perubahan hanya di memori")
            self._notify()
            if self.last_save_error:
                print(f"[WARNING] followup: {self.last_save_error}")
            return False
        snap = {k: _pyval(row.get(k, "")) for k in SNAPSHOT_FIELDS}
        snap["fp"] = fp
        snap["added_at"] = datetime.now().isoformat(timespec="seconds")
        self.rows[fp] = snap
        ok = self.save()
        self.last_save_error = None if ok else (
            "gagal menyimpan status followup (file terkunci/disk?) "
            "— perubahan hanya di memori")
        self._notify()
        if self.last_save_error:
            print(f"[WARNING] followup: {self.last_save_error}")
        return True

    def clear(self) -> None:
        if not self.rows:
            return
        self.rows.clear()
        self.resolved.clear()
        # v3.9.42 — audit 61-c/61-d (F5 P3): pola toggle — hasil save tidak
        # diabaikan (gagal persist → [WARNING], lihat toggle).
        # v3.9.59 — FIX 2f-08: + flag last_save_error utk view (lihat toggle).
        ok = self.save()
        ok_r = self.save_resolved()
        self.last_save_error = None if (ok and ok_r) else (
            "gagal menyimpan status followup (file terkunci/disk?) "
            "— perubahan hanya di memori")
        self._notify()
        if self.last_save_error:
            print(f"[WARNING] followup: {self.last_save_error}")

    # -- persistensi ---------------------------------------------------------------
    def save(self) -> bool:
        return _atomic_save(followup_path(), {
            "version": 1,
            "rows": {fp: r for fp, r in self.rows.items()},
        })

    def save_resolved(self) -> bool:
        return _atomic_save(resolved_path(), {
            "version": 1, "rows": self.resolved,
        })

    def load(self) -> None:
        p = followup_path()
        if p.exists():
            try:
                data = json.loads(p.read_text(encoding="utf-8"))
                rows = data.get("rows", {}) if isinstance(data, dict) else {}
                if isinstance(rows, dict):
                    self.rows = {
                        str(k): r for k, r in rows.items()
                        if isinstance(r, dict)
                    }
            except (json.JSONDecodeError, ValueError, OSError) as e:
                _LOG.warning("followup.json korup (%s) — mulai kosong", e)
                _backup_corrupt(p)
                self.rows = {}
                # v3.9.59 — FIX 2f-07: pulihan kini TERLIHAT (flag +
                # [WARNING] log internal + dialog bila GUI — dulu senyap,
                # user kehilangan seluruh daftar tanpa pesan).
                _notify_recovery(self, p, str(e))
        pr = resolved_path()
        if pr.exists():
            try:
                data = json.loads(pr.read_text(encoding="utf-8"))
                rows = data.get("rows", {})
                if isinstance(rows, dict):
                    self.resolved = {str(k): r for k, r in rows.items()
                                     if isinstance(r, dict)}
            except Exception as e:
                _LOG.warning("followup_resolved.json korup → backup .bad")
                _backup_corrupt(pr)
                self.resolved = {}
                # v3.9.59 — FIX 2f-07: mirror jalur followup.json — status
                # TL/PIC yang hilang juga diberitahukan (flag + log).
                _notify_recovery(self, pr, f"korup: {e}")

    # -- export ------------------------------------------------------------------
    def export_xlsx(self, path: str) -> str:
        """Tulis sheet 'FOLLOW UP' — FORMAT FILE SAMPLE (v3.9.6).

        Urutan kolom persis sheet Follow UP workbook asli user:
        Resto · Nama Barang · [SO tgl H-2..H+2 — heading angka tanggal,
        mis. H=25 → 23/24/25/26/27] · Sistem · Deviasi · Loss Surplus to
        BOM · Keterangan · SO · DEVIASI · Selisih · Selisih Deviasi ·
        Cek FU · PIC (+ kolom pelacakan app: Asal · Status TL ·
        Ditambahkan). Baris resolved membawa nilai kerja FU tersimpan;
        Deviasi/SO per tanggal negatif diberi font merah. Freeze +
        autofilter.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        # heading kolom tanggal = angka tanggal (dari H snapshot terbaru)
        h_ref = None
        for snap in reversed(list(self.rows.values())):
            h_val = snap.get("h")
            if isinstance(h_val, (int, float)) and h_val:
                h_ref = int(h_val)
                break
        if h_ref is None:
            header = list(EXPORT_COLS)
        else:
            header = [f"{h_ref + DATE_COL_OFFSETS[c]}" if c in DATE_COL_IDS
                      else c for c in EXPORT_COLS]

        wb = Workbook()
        ws = wb.active
        ws.title = "FOLLOW UP"
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # posisi kolom (indeks 0-based) utk pewarnaan
        idx = {c: header.index(c) for c in header if c in header}
        dev_col = idx.get("Deviasi")
        dev2_col = idx.get("DEVIASI")
        date_cols = [j for j, c in enumerate(header)
                     if c not in ("Resto", "Nama Barang", "Sistem",
                                  "Deviasi", "Loss Surplus to BOM",
                                  "Keterangan", "DEVIASI", "Selisih",
                                  "Selisih Deviasi", "Cek FU", "PIC",
                                  "Asal", "Status TL", "Ditambahkan")]

        order = list(self.rows.values())
        n_handover = len(HANDOVER_COLS)
        handover_start = len(header) - n_handover - 3  # 3 kolom pelacakan app
        for i, snap in enumerate(order):
            r = i + 2
            info = self.resolved.get(snap.get("fp", ""), {})
            vals = [
                snap.get("resto", ""), snap.get("barang", ""),
                snap.get("so_m2", ""), snap.get("so_m1", ""),
                snap.get("so", ""), snap.get("so_p1", ""),
                snap.get("so_p2", ""),
                snap.get("sistem", ""), snap.get("deviasi", ""),
                snap.get("loss_bom", ""), snap.get("cek", ""),
                snap.get("so", ""), snap.get("deviasi", ""),
                info.get("Selisih", ""), info.get("Selisih Deviasi", ""),
                info.get("Cek FU", ""), info.get("PIC", ""),
                snap.get("asal", ""),
                "Selesai" if info else "Menunggu",
                _fmt_added(snap.get("added_at", "")),
            ]
            for j, v in enumerate(vals):
                cell = ws.cell(row=r, column=j + 1, value=v)
                neg = (j in date_cols or j in (dev_col, dev2_col)) \
                    and isinstance(v, (int, float)) and v < 0
                if neg:
                    cell.font = Font(color="DC2626")
            # resolved → fill hijau lembut di kolom kerja FU + PIC
            if info:
                fill = PatternFill("solid", start_color="ECFDF5")
                for j in range(handover_start, handover_start + n_handover):
                    ws.cell(row=r, column=j + 1).fill = fill

        ws.freeze_panes = "A2"
        if order:
            ws.auto_filter.ref = (
                f"A1:{get_column_letter(len(header))}{len(order) + 1}")
        widths = {"Resto": 10, "Nama Barang": 30, "Sistem": 11,
                  "Deviasi": 11, "DEVIASI": 11, "SO": 10,
                  "Loss Surplus to BOM": 12, "Keterangan": 18,
                  "Selisih": 9, "Selisih Deviasi": 12, "Cek FU": 8,
                  "PIC": 14, "Asal": 12, "Status TL": 10,
                  "Ditambahkan": 16}
        for j, c in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(j)].width = \
                widths.get(c, 8 if c.isdigit() else 12)
        # v3.9.59 — FIX 2f-04 (audit 2-f): tulis ATOMIK (tmp + os.replace —
        # pola 6a-02 app._export_excel). wb.save() dulu menulis LANGSUNG
        # ke path tujuan → exception di tengah (disk penuh / styling)
        # meninggalkan .xlsx PARSIAL yang MENIMPA export LAMA yang baik
        # di path sama. Kini file tujuan hanya tersentuh os.replace()
        # atomik SETELAH sukses penuh; kegagalan apa pun (finally)
        # menghapus tmp (tmp ber-ekstensi .xlsx).
        tmp_path = str(path) + ".tmp.xlsx"
        _ok = False
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, str(path))
            _ok = True
        finally:
            if not _ok:
                try:
                    if os.path.exists(tmp_path):
                        os.unlink(tmp_path)
                except OSError:
                    pass
        return path


