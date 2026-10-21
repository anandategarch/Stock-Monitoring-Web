"""
helpers.py — Helper functions yang mirror persis fnBukaSheet, fnCombineFolder,
dan fungsi text delimiter di Power Query M code.

OPTIMIZED untuk speed: vectorized text operations via .str accessor.

v3.9.37 — CACHE BACAAN FILE SUMBER (user: "ada tombol refresh untuk file
sumber, bisa gak? dan gak harus refresh semua tapi"): fn_buka_sheet kini
menyimpan bacaan per (file, sheet, header_row, raw) divalidasi SIGNATURE
(mtime_ns + ukuran). File yang tak berubah dibaca SEKALI per sesi —
run berikutnya (termasuk Refresh Terpilih) menyajikan frame dari cache
(copy defensif — pemanggil bebas memutasi). Invalidasi selektif via
read_cache_invalidate(paths) (dipakai dialog 🔄 Refresh File Sumber);
hasil identik baca ulang — hanya lebih cepat. Batas memori: estimasi
sel object (guard _READ_CACHE_MAX_CELLS, evict FIFO).

v3.9.42 — AUDIT FIX (61-a):
- F2: to_datetime_robust — string tanggal kini dayfirst=True + format='mixed'
  (konteks Indonesia dd/mm/yyyy: '01/10/2026' dulu jadi 10 Jan, '13/10/2026'
  dulu NaT — hari salah/tanggal hilang senyap; jalur serial Excel numerik
  tetap apa adanya; 'mixed' mencegah inferensi pandas membalik kolom ISO
  jadi %Y-%d-%m).
- F4: fn_combine_folder — kontribusi baris per file kini DI-LOG
  ([INFO] rincian + [WARNING] bila >1 file menyumbang baris — folder
  berisi export lama+baru/salinan menjumlah SALDO/IN/OUT 2× SENYAP
  sebelumnya).
"""
import glob
import os
import re
import threading
from typing import List, Optional

import numpy as np
import pandas as pd

# =============================================================================
# TEXT DELIMITER FUNCTIONS (vectorized via .str accessor)
# =============================================================================

def text_before_delimiter(s, delimiter: str = ".") -> Optional[str]:
    """Mirror M: Text.BeforeDelimiter(text, delimiter). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    idx = s.find(delimiter)
    if idx < 0:
        return s
    return s[:idx]


def text_after_delimiter(s, delimiter: str = ".") -> Optional[str]:
    """Mirror M: Text.AfterDelimiter(text, delimiter). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    idx = s.find(delimiter)
    if idx < 0:
        return s
    return s[idx + len(delimiter):]


def text_between_delimiters(s, open_delim: str = "(", close_delim: str = ")") -> Optional[str]:
    """Mirror M: Text.BetweenDelimiters(text, openDelim, closeDelim). Single value."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return None
    s = str(s)
    start = s.find(open_delim)
    if start < 0:
        return None
    start += len(open_delim)
    end = s.find(close_delim, start)
    if end < 0:
        return None
    return s[start:end]


# === VECTORIZED VERSIONS (for pandas Series — 10-50x faster than .apply) ===

def vec_text_before_delimiter(series: pd.Series, delimiter: str = ".") -> pd.Series:
    """Vectorized Text.BeforeDelimiter untuk pandas Series.

    Perilaku: baris mengandung delimiter → teks sebelum delimiter;
    tanpa delimiter → teks asli; NaN → None. (v3.3: dead lines dibersihkan
    — dulu `result` dihitung dua kali, hasil pertama ditimpa.)
    """
    s = series.astype(str).where(series.notna(), other=np.nan)
    idx = s.str.find(delimiter)
    result = s.copy()
    has_delim = idx >= 0
    if bool(has_delim.any()):
        positions = idx[has_delim].values
        vals = s[has_delim].values
        result[has_delim] = [str(v)[:int(p)] for v, p in zip(vals, positions)]
    result = result.where(series.notna(), other=None)
    return result


def vec_text_after_delimiter(series: pd.Series, delimiter: str = ".") -> pd.Series:
    """Vectorized Text.AfterDelimiter untuk pandas Series."""
    s = series.astype(str)
    parts = s.str.split(delimiter, n=1, expand=True)
    if parts.shape[1] < 2:
        # No delimiter found in any row
        return series
    result = parts[1].fillna(series.astype(str))
    result = result.where(series.notna(), other=None)
    return result


def vec_text_between_delimiters(series: pd.Series, open_delim: str = "(", close_delim: str = ")") -> pd.Series:
    """Vectorized Text.BetweenDelimiters untuk pandas Series.
    Uses regex: extract text between first open_delim and next close_delim.
    """
    # Escape regex special chars in delimiters
    open_esc = re.escape(open_delim)
    close_esc = re.escape(close_delim)
    pattern = f"{open_esc}([^{close_esc}]*){close_esc}"
    result = series.astype(str).str.extract(pattern, expand=False)
    # Where no match, result is NaN → convert to None
    result = result.where(result.notna(), other=None)
    # Also preserve original NaN
    result = result.where(series.notna(), other=None)
    return result


def kategori_blank_mask(kat: "pd.Series | None") -> "pd.Series | None":
    """v3.9.44 — mask baris TANPA kategori master (NaN/blank/'nan'/'none').

    AKAR "BUMBU PASTA KUAH MLGPAR masih belum ada di tab Deviasi GR"
    (probe E2E data nyata v3.9.44): item 42.17-only kode 100383 lolos
    whitelist versi-toleran (v3.9.41) dan masuk pipeline sampai pivot,
    TAPI tidak terdaftar di master '21. DBS Kategori' user (bukti
    workbook: hanya varian (V.20) ada di 'Data Base Aktif' & 'Data Base
    Cek Deviasi') → KATEGORI = NaN → filter KATEGORI default 'Deviasi'
    di view menyembunyikan barisnya ('nan' ≠ 'DEVIASI') — item lenyap
    DARI VIEW padahal ADA di data (hanya terlihat bila user ganti ke
    'Semua' — tanpa tahu itu, keluhan "masih belum ada" muncul lagi).

    KEBIJAKAN v3.9.45 (revisi dari v3.9.44): baris yang master-nya
    tidak kenal TIDAK BOLEH disembunyikan oleh seleksi DEFAULT 'Deviasi'
    — kelas bug "hilang senyap karena master kurang entri". v3.9.44
    sempat memasang passthrough di seleksi APA PUN, tapi dengan master
    yang men-cover sebagian barang (90% baris '-' di data nyata) filter
    kategori jadi no-op (user: "ganti tidak terjadi apa apa") → kini
    passthrough HANYA seleksi default (lihat kategori_passthrough) +
    entry combobox KATEGORI_BLANK_LABEL utk akses eksplisit; seleksi
    lain strict. Nilai kategorinya tetap kosong (jujur — tidak
    memfabrikasi 'DEVIASI').
    Konsumen: widgets (re-export) → deviasi_view · nominal_resto_view ·
    pivot_view (2 titik) + append1 jejak saldo-only tanpa kategori.
    """
    if kat is None:
        return None
    s = kat.astype(str).str.strip()
    # '-' = konvensi DISPLAY kategori kosong (deviasi_view kategori_s
    # fillna('-') — pola in_selisih_view); 'nan'/'none' = artefak
    # astype(str)/safe_to_str dari NaN master.
    blank = s.str.lower().isin(["", "nan", "none", "-", "—"])
    try:
        blank = kat.isna() | blank
    except (TypeError, ValueError):
        pass  # series non-standar — mask string sudah cukup
    return blank


# v3.9.45 — label combobox utk memilih baris TANPA kategori master
# secara EKSPLISIT (akses kelas item 100383 tanpa menonaktifkan filter
# kategori barang lain).
KATEGORI_BLANK_LABEL = "(Tanpa Kategori)"

# v3.9.49 — KATEGORI DEFAULT GABUNGAN (user: "sebenernya sudah ada
# KATEGORI dari file Rekap Mentah folder Deviasi Weekly jadi bisa ambil
# kategori dari situ jadi secara default di semua tab kategori sudah
# pilih: COM DEVIASI - RESTO dan Biaya Packaging - RESTO"): kolom
# KATEGORI file Rekap Mentah (sheet '1  DEVIASI WEEKLY') kini DIBAWA
# pipeline (fn_rekap_mentah → query_full_rekap_mentah_so → append1
# override PRIORITAS atas master '21. DBS Kategori' — barang di rekap
# memakai kategori rekapnya; barang hanya-di-master/in-only tetap
# memakai master). Dua nilai kategori rekap user dijadikan seleksi
# DEFAULT gabungan di semua tab ber-filter kategori (Deviasi GR/PCS ·
# 🧮 Nominal per Resto · Pivot · default header Rekap Selisih IN).
# v3.9.55 — LABEL DIPENDEKAN (user: "Kategori kembali ke DBS kategori
# barang aja lebih valid atau tambah kategorinya jadi 2 aja dan default
# Deviasi"): entry gabungan tetap memilih TEPAT 2 kategori rekap (COM
# DEVIASI - RESTO + Biaya Packaging - RESTO — semantics v3.9.49/53 utuh)
# TAPI labelnya kini RINGKA "Deviasi (COM + Packaging)" — dropdown
# kembali tampil sebagai daftar kategori DBS (label virtual 49-char
# panjang dianggap kurang valid oleh user). Semantik STRICT v3.9.53
# (tanpa passthrough blank) TIDAK berubah; label lama 49-char tetap
# dikenali kategori_multi_values (kompatibilitas seleksi tersimpan).
KATEGORI_REKAP_VALUES = ("COM DEVIASI - RESTO", "Biaya Packaging - RESTO")
KATEGORI_DEFAULT_LABEL = "Deviasi (COM + Packaging)"
# v3.9.55 — label panjang v3.9.49-54 tetap cocok (seleksi tersimpan /
# pemanggil programatik lama tak berubah perilakunya).
_KATEGORI_DEFAULT_LABEL_LEGACY = "COM DEVIASI - RESTO + Biaya Packaging - RESTO"


def kategori_multi_values(sel) -> "tuple[str, ...] | None":
    """v3.9.49 — himpunan kategori yang dicocokkan seleksi GABUNGAN
    (entry default KATEGORI_DEFAULT_LABEL), atau None bila seleksi bukan
    entri gabungan.

    Konsumen memakai ini SEBELUM cabang kategori_passthrough. v3.9.53 —
    seleksi gabungan kini STRICT: view memfilter HANYA nilai gabungan
    (passthrough baris kosong DIHAPUS — lihat kategori_passthrough (d);
    user menemukan barang tanpa kategori BASKOM dll. ikut tampil).
    v3.9.50 — [audit 3-b] case-INSENSITIVE (strip+lower — parity
    kategori_passthrough; dulu exact-case: pemanggil programatik dgn
    kapitalisasi beda jatuh ke cabang strict → 0 baris + blank-pass
    tanpa disengaja).
    """
    if sel is None:
        return None
    _s = str(sel).strip().lower()
    # v3.9.55 — label baru RINGKA + label lama 49-char (legacy) dua-duanya
    # dikenali → entri gabungan yang sama.
    if _s == KATEGORI_DEFAULT_LABEL.lower() \
            or _s == _KATEGORI_DEFAULT_LABEL_LEGACY.lower():
        return KATEGORI_REKAP_VALUES
    return None


def kategori_filter_mask(kat, sel):
    """v3.9.49 — mask baris yang cocok seleksi kategori GABUNGAN
    (case-insensitive + strip — entri gabungan VIRTUAL, bukan nilai data
    tunggal). None bila seleksi bukan entri gabungan atau kat None."""
    vals = kategori_multi_values(sel)
    if vals is None or kat is None:
        return None
    s = kat.astype(str).str.strip().str.upper()
    return s.isin([v.upper() for v in vals])


def kategori_match_mask(kat, sel):
    """v3.9.54 — mask KANONIK seleksi kategori: SEMUA cabang entri
    VIRTUAL dalam satu fungsi (semantik PERSIS jalur tampilan view —
    pivot_view _name_values/_compute_payload · nominal_resto_view ·
    deviasi_view).

    AKAR BUG (user: "ADA BUG di pop up pivot tabel yang kosong ketika aku
    double klik"): _on_double_click pivot meneruskan seleksi KATEGORI
    combobox (DEFAULT = entri gabungan virtual KATEGORI_DEFAULT_LABEL)
    ke show_details_popup yang memfilter EXACT string — label virtual
    bukan nilai data → 0 baris → popup KOSONG pada setiap double-click
    dengan kategori default sejak v3.9.49.

    Cabang (urutan & semantik identik blok filter view):
    1. KATEGORI_BLANK_LABEL → kategori_blank_mask (baris tanpa kategori);
    2. entri gabungan → kategori_filter_mask (isin multi-nilai virtual,
       case-insensitive);
    3. legacy 'Deviasi' (kategori_passthrough) → exact | blank;
    4. seleksi eksplisit lain → exact string match.

    Return mask boolean Series; None bila `kat` None ATAU `sel` bukan
    filter aktif (None/kosong/'(All)') — pemanggil melanjutkan tanpa
    memfilter.
    """
    if kat is None or sel is None:
        return None
    sel_s = str(sel)
    if not sel_s.strip() or sel_s == "(All)":
        return None
    if sel_s == KATEGORI_BLANK_LABEL:
        return kategori_blank_mask(kat)
    m = kategori_filter_mask(kat, sel_s)
    if m is not None:
        return m
    if kategori_passthrough(sel_s):
        blank = kategori_blank_mask(kat)
        return (kat.astype(str) == sel_s) | blank
    return kat.astype(str) == sel_s


def kategori_has_default_values(series) -> bool:
    """v3.9.49 — apakah kolom KATEGORI memuat salah satu nilai rekap
    default (dipakai view utk memasang entry gabungan + seleksi default
    hanya bila relevan dgn data — tanpa itu, entry default di data tanpa
    kategori rekap akan memfilter 0 baris secara membingungkan)."""
    if series is None:
        return False
    try:
        s = series.dropna().astype(str).str.strip().str.upper()
    except (AttributeError, TypeError, ValueError):
        return False
    return bool(s.isin([v.upper() for v in KATEGORI_REKAP_VALUES]).any())


def kategori_passthrough(sel) -> bool:
    """v3.9.45 — apakah seleksi kategori meng-include baris TANPA
    kategori master (passthrough kategori_blank_mask)?

    HANYA seleksi DEFAULT 'Deviasi' (case-insensitive). Latar v3.9.45
    (user: "filter kategori di tab deviasi gr dan pcs kok seperti
    ngebug? aku ganti tidak terjadi apa apa"): v3.9.44 memasang
    passthrough di SEMUA seleksi — dengan master '21. DBS Kategori'
    yang hanya men-cover SEBAGIAN barang (probe E2E data nyata: 90%
    baris pivot ber-KATEGORI '-'), baris '-' tampil di seleksi APA PUN
    → mengganti kategori praktis tidak mengubah apa pun (Semua ==
    DEVIASI == 1.347 baris di probe) — filter jadi no-op.

    KEBIJAKAN (v3.9.45 → direvisi v3.9.53):
    (a) default 'Deviasi' TETAP meng-include baris tanpa kategori —
        perlindungan kelas 100383 era pra-kategori-rekap (barang
        ber-deviasi yang belum masuk master tidak boleh lenyap senyap
        dari tab Deviasi — keluhan v3.9.44);
    (b) seleksi eksplisit LAIN = STRICT (filter bekerja normal —
        keluhan v3.9.45);
    (c) baris tanpa kategori punya entry combobox sendiri
        (KATEGORI_BLANK_LABEL) utk akses eksplisit;
    (d) v3.9.53 — entri GABUNGAN default KATEGORI_DEFAULT_LABEL kini
        STRICT (TIDAK ber-passthrough): user: "kok ada nama barang yang
        bukan kategori COM Deviasi + Packaging Deviasi misal BASKOM
        STAINLESS 18 CM, BASKOM STAINLESS UK 26 CM dll coba cek apa ini
        bug?" — passthrough blank membuat barang TANPA kategori (mis.
        perlengkapan BASKOM) ikut tampil di seleksi default. Proteksi
        kelas 100383 kini DIPENUHI KATEGORI REKAP v3.9.49 (barang
        ber-deviasi ber-kategori rekap tetap tampil di seleksi default,
        mis. BUMBU PASTA KUAH (V.20)); kelas tanpa kategori tetap
        terjangkau lewat entry '(Tanpa Kategori)' + jumlah tersembunyi
        dinyatakan di status view — tidak lenyap senyap.
    Konsumen: deviasi_view ×2 · nominal_resto_view · pivot_view ×2
    (via re-export widgets).
    """
    if sel is None:
        return False
    return str(sel).strip().lower() == "deviasi"


# =============================================================================
# COMBINE FOLDER (mirror fnCombineFolder + fnBukaSheet)
# =============================================================================

def _list_excel_files(folder_path: str) -> List[str]:
    """List semua file Excel di folder, exclude hidden & temp files.
    v3.3.1: sertakan *.xlsb — sumber SO user kini berupa workbook gabungan
    .xlsb ("01.1 CEK STOCK V2.xlsb") yang sebelumnya TIDAK terbaca sama
    sekali (akibatnya Saldo SO kosong semua di tab Deviasi GR/PCS)."""
    if not os.path.isdir(folder_path):
        return []
    patterns = ["*.xlsx", "*.xls", "*.xlsm", "*.xlsb"]
    files = []
    for p in patterns:
        files.extend(glob.glob(os.path.join(folder_path, p)))
    # Filter: exclude hidden files (check basename, not full path) & temp lock files (~$...)
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    files = [f for f in files if not os.path.basename(f).startswith(".")]
    return sorted(files)


def _get_excel_engine() -> str:
    """Detect best available Excel engine.
    Priority: calamine (Rust, 10-20x faster) > openpyxl (default).
    """
    try:
        import python_calamine  # noqa: F401
        return "calamine"
    except ImportError:
        return None  # default to openpyxl


# Cache engine detection (only check once)
_EXCEL_ENGINE = _get_excel_engine()


# =============================================================================
# v3.9.37 — CACHE BACAAN FILE SUMBER (refresh selektif tanpa baca ulang
# file yang tak berubah; user: "gak harus refresh semua")
# =============================================================================
_READ_CACHE: dict = {}          # key -> {"sig", "df", "cells"}
_READ_CACHE_LOCK = threading.Lock()
_READ_CACHE_CELLS = 0           # total sel ter-cache (guard memori)
_READ_CACHE_MAX_CELLS = 20_000_000   # ~estimasi 20 juta sel object
_READ_CACHE_HIT_FILES: set = set()   # normpath file yang disajikan cache
_READ_CACHE_MISS_FILES: set = set()  # normpath file yang dibaca disk


def _read_cache_key(file_path: str, sheet_name, header_row: int,
                    raw: bool):
    """Kunci cache: path ternormalisasi (Windows case-insensitive) +
    sheet + header_row + raw — signature file divalidasi TERPISAH."""
    return (os.path.normcase(os.path.normpath(file_path)),
            str(sheet_name), int(header_row), bool(raw))


def _read_cache_sig(file_path: str):
    """Signature file: (mtime_ns, ukuran). None = file tak terbaca stat
    (hilang/terkunci) → cache MISS (baca ulang; kegagalan mengikuti jalur
    error fn_buka_sheet yang sudah ada)."""
    try:
        st = os.stat(file_path)
        return (st.st_mtime_ns, st.st_size)
    except OSError:
        return None


def _read_cache_get(key, sig):
    """Ambil frame dari cache bila signature cocok → COPY defensif
    (pemanggil boleh memutasi hasilnya — insert kolom Name / ffill di
    fn_combine_folder — tanpa menodai entry cache)."""
    with _READ_CACHE_LOCK:
        entry = _READ_CACHE.get(key)
        if entry is None or sig is None or entry["sig"] != sig:
            return None
        _READ_CACHE_HIT_FILES.add(key[0])
        return entry["df"].copy()


def _read_cache_put(key, sig, df) -> None:
    """Simpan copy frame ke cache (evict FIFO bila guard sel terlampaui).
    Frame yang DISIMPAN adalah copy — frame asli dikembalikan ke
    pemanggil (boleh dimutasi bebas). Frame yang melebihi batas sel
    sendirian DITOLAK sebelum eviksi (cache lain tak diganggu)."""
    if sig is None or df is None:
        return
    try:
        cells = int(df.shape[0]) * int(df.shape[1])
    except (TypeError, ValueError):
        return
    if cells <= 0:
        return
    global _READ_CACHE_CELLS
    with _READ_CACHE_LOCK:
        # v3.9.70 — audit ulang v3.9.69 (P2): frame raksasa yang melebihi
        # guard SENDIRIAN ditolak SEBELUM eviksi — dulu loop eviksi
        # mengosongkan SELURUH cache demi frame yang akhirnya ditolak
        # guard juga (cache lain terbuang percuma).
        if cells > _READ_CACHE_MAX_CELLS:
            return
        # v3.9.70 — audit ulang v3.9.69 (P2, lanjutan fix v3.9.69): entry
        # lama key-sama DILEPAS dari dict (pop) → kontribusi selnya
        # dikembalikan SEKALI. Fix v3.9.69 hanya mengurangi counter TANPA
        # melepas entry → bila entry lama ikut ter-eviksi FIFO, selnya
        # dikurangi DUA KALI (contoh audit, batas 100: A=30 B=40 C=20 →
        # ganti A→50 → eviksi hapus A-lama → counter 80 padahal isi dict
        # 110; counter bisa negatif → keputusan eviksi/tolak kacau).
        # Salah akuntansi, bukan salah data. (Riwayat v3.9.69/P2.1:
        # dulu counter hanya ditambah sel baru — 100 → ditimpa 150 →
        # counter 250; pemicu nyata sig-mismatch mtime+size antar-run.)
        _old = _READ_CACHE.pop(key, None)
        if _old is not None:
            _READ_CACHE_CELLS -= _old.get("cells", 0)
        # evict FIFO sampai muat (dict = urutan insert)
        while (_READ_CACHE_CELLS + cells > _READ_CACHE_MAX_CELLS
               and _READ_CACHE):
            _old_k, _old = next(iter(_READ_CACHE.items()))
            _READ_CACHE.pop(_old_k, None)
            _READ_CACHE_CELLS -= _old.get("cells", 0)
        if _READ_CACHE_CELLS + cells > _READ_CACHE_MAX_CELLS:
            return  # defensif (hanya bila MAX menyusut saat runtime)
        _READ_CACHE[key] = {"sig": sig, "df": df.copy(), "cells": cells}
        _READ_CACHE_CELLS += cells
        _READ_CACHE_MISS_FILES.add(key[0])


def read_cache_invalidate(paths=None) -> int:
    """v3.9.37 — buang entry cache file sumber. Return jumlah entry
    yang dibuang.

    paths=None → SEMUA entry dibuang ("Refresh Semua" = baca ulang
    semua file dari disk). Bila list/set path → hanya entry yang
    file-nya termasuk (pencocokan normcase+normpath — jalur Windows
    beda kapital / separator tetap kena)."""
    global _READ_CACHE_CELLS
    with _READ_CACHE_LOCK:
        if paths is None:
            n = len(_READ_CACHE)
            _READ_CACHE.clear()
            _READ_CACHE_CELLS = 0
            return n
        targets = {os.path.normcase(os.path.normpath(str(p)))
                   for p in paths}
        drop = [k for k in _READ_CACHE if k[0] in targets]
        for k in drop:
            _READ_CACHE_CELLS -= _READ_CACHE[k].get("cells", 0)
            _READ_CACHE.pop(k, None)
        return len(drop)


def read_cache_stats() -> tuple:
    """v3.9.37 — (n_entry, n_file_unik, n_file_hit, n_file_miss) utk
    log observabilitas (file ter-cache / disajikan cache / dibaca disk)."""
    with _READ_CACHE_LOCK:
        files = {k[0] for k in _READ_CACHE}
        return (len(_READ_CACHE), len(files),
                len(_READ_CACHE_HIT_FILES), len(_READ_CACHE_MISS_FILES))


def read_cache_reset_counters() -> None:
    """v3.9.37 — reset penghitung hit/miss per-file (dipanggil app di
    AWAL tiap run pipeline supaya log akhir run mencerminkan run itu)."""
    with _READ_CACHE_LOCK:
        _READ_CACHE_HIT_FILES.clear()
        _READ_CACHE_MISS_FILES.clear()


def _xlsb_engine() -> Optional[str]:
    """Engine untuk file .xlsb: calamine (jika ada) > pyxlsb > None.
    None = engine tidak tersedia (.xlsb harus di-skip dengan pesan jelas)."""
    if _EXCEL_ENGINE == "calamine":
        return "calamine"
    try:
        import pyxlsb  # noqa: F401
        return "pyxlsb"
    except ImportError:
        return None


def _list_sheet_names(file_path: str) -> List[str]:
    """Daftar nama sheet dari file Excel (semua format, termasuk .xlsb).
    Urutan coba: calamine > openpyxl (.xlsx) / pyxlsb (.xlsb)."""
    low = file_path.lower()
    # 1) calamine — cepat untuk semua format
    if _EXCEL_ENGINE == "calamine":
        try:
            from python_calamine import load_workbook
            return list(load_workbook(file_path).sheet_names)
        except Exception:
            pass
    # 2) .xlsb → pyxlsb
    if low.endswith(".xlsb"):
        try:
            from pyxlsb import open_workbook
            with open_workbook(file_path) as wb:
                return list(wb.sheets)
        except Exception:
            return []
    # 3) .xlsx/.xlsm → openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception:
        return []


def fn_buka_sheet(file_path: str, sheet_name: str, header_row: int = 0,
                   raw: bool = False) -> pd.DataFrame:
    """
    Mirror M: fnBukaSheet(binary, sheetName, doPromote).
    Baca sheet tertentu dari file Excel.

    OPTIMIZED: gunakan calamine engine (Rust) jika tersedia — 10-20x faster.
    Fallback ke openpyxl jika calamine tidak terinstall.
    v3.3.1: file .xlsb didukung (engine pyxlsb bila calamine absen).

    raw=True: baca tanpa header (header=None), kolom = Column1, Column2, ...
              Mirror M doPromote=null/false.
    raw=False: baca dengan header di baris header_row (0-indexed).
              Mirror M doPromote=true + Skip(N) + PromoteHeaders.
              header_row=4 = skip rows 0-3, row 4 jadi header.
    """
    # v3.3.1: pilih engine eksplisit untuk .xlsb (openpyxl tidak bisa .xlsb;
    # pandas default-engine pun butuh pyxlsb terinstall — beri pesan jelas)
    # v3.9.37: cache bacaan (signature mtime+ukuran) — file tak berubah
    # tak dibaca ulang dari disk (Refresh Terpilih: hanya file pilihan
    # user yang dibaca ulang; hasil identik).
    _rc_key = _read_cache_key(file_path, sheet_name, header_row, raw)
    _rc_sig = _read_cache_sig(file_path)
    _rc_hit = _read_cache_get(_rc_key, _rc_sig)
    if _rc_hit is not None:
        return _rc_hit
    if file_path.lower().endswith(".xlsb") and _EXCEL_ENGINE != "calamine":
        xlsb_eng = _xlsb_engine()
        if xlsb_eng is None:
            raise ImportError(
                "File .xlsb membutuhkan package 'pyxlsb' atau "
                "'python-calamine' — jalankan: pip install pyxlsb")
        engine_kwargs = {"engine": xlsb_eng}
    else:
        engine_kwargs = {"engine": _EXCEL_ENGINE} if _EXCEL_ENGINE else {}
    try:
        if raw:
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None,
                                dtype=object, **engine_kwargs)
        else:
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row,
                                dtype=object, **engine_kwargs)
        _read_cache_put(_rc_key, _rc_sig, df)
        return df
    except Exception as e:
        # Fallback to openpyxl if calamine fails (e.g., unsupported format)
        if _EXCEL_ENGINE:
            if raw:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None, dtype=object)
            else:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row, dtype=object)
            _read_cache_put(_rc_key, _rc_sig, df)
            return df
        raise e


def fn_combine_folder(folder_path: str, sheet_name: str, header_row: int = 0,
                       raw: bool = False, add_name_col: bool = False,
                       ffill_cols: "list[str] | None" = None) -> pd.DataFrame:
    """
    Mirror M: fnCombineFolder(folderPath, sheetName, doPromote).

    raw=True: baca tanpa header per file, gabungkan raw.
              Mirror M doPromote=null/false.
    raw=False: baca dengan header_row per file, gabungkan by column name.
              Mirror M doPromote=true.

    add_name_col: jika True, tambah kolom "Name" berisi nama file (basename).
                  Diperlukan untuk fnRekapMentah (NAMA RESTO dari file name).

    ffill_cols: v3.9.33 — AUDIT FIX (48-a): list nama kolom yang di-ffill
                PER FILE sebelum concat. Nilai ber-pola grup header (mis.
                kolom 'Kategori Barang' 42.17 — hanya terisi di baris
                pertama grup) harus di-ffill DALAM file yang sama; ffill
                pada frame hasil CONCAT membuat file berikutnya mewarisi
                nilai EKOR file sebelumnya (kategori lintas file bleed —
                uji 48-a: file-2 grup pertama mewarisi kategori terakhir
                file-1). Bisa berisi beberapa kandidat nama kolom — semua
                yang ada di file di-ffill. Default None = perilaku lama.
    """
    files = _list_excel_files(folder_path)
    if not files:
        return pd.DataFrame()

    dfs = []
    # v3.9.42 — AUDIT FIX (61-a, F4): kontribusi baris per file dikumpulkan
    # SEBELUM concat (mirror log fn_9901 v3.9.33) — concat dulu SENYAP:
    # folder berisi 2 file (export lama + baru / salinan) menjumlahkan
    # SALDO/IN/OUT 2× tanpa jejak apa pun di log.
    kontribusi = []      # (basename, n_baris) per file yang berhasil dibaca
    for f in files:
        try:
            if raw:
                df = fn_buka_sheet(f, sheet_name, raw=True)
            else:
                df = fn_buka_sheet(f, sheet_name, header_row=header_row)
            if df is not None and len(df) > 0:
                if add_name_col:
                    # Insert Name column at position 0
                    df.insert(0, "Name", os.path.basename(f))
                if ffill_cols:
                    # v3.9.33 — AUDIT FIX (48-a): ffill PER FILE sebelum
                    # concat (mirror pola per-file fn_4109) — grup kategori
                    # file berikutnya TIDAK mewarisi ekor file sebelumnya.
                    for c in ffill_cols:
                        if c in df.columns:
                            df[c] = df[c].ffill()
                dfs.append(df)
                kontribusi.append((os.path.basename(f), len(df)))
        except Exception as e:
            print(f"  [WARNING] Skip file {os.path.basename(f)}: {e}")

    if not dfs:
        return pd.DataFrame()

    if raw:
        # Raw mode: concat by position (columns are 0, 1, 2, ...)
        # Use ignore_index + sort=False for speed
        combined = pd.concat(dfs, ignore_index=True, sort=False)
    else:
        # Header mode: concat by column name
        combined = pd.concat(dfs, ignore_index=True, sort=False)

    # v3.9.42 — AUDIT FIX (61-a, F4): log kontribusi per file — [INFO]
    # ringkas selalu tercetak (observabilitas); >1 file menyumbang baris
    # → [WARNING] mencolok (pola log duplikat fn_9901 v3.9.33): file
    # kembar/salinan export di folder data menjumlahkan baris 2×.
    if kontribusi:
        rincian = ", ".join(f"{nama}={n}" for nama, n in kontribusi)
        print(f"  [INFO] combine_folder {folder_path}: {len(combined)} baris "
              f"dari {len(kontribusi)} file ({rincian})")
        if len(kontribusi) > 1:
            print(f"  [WARNING] combine_folder {folder_path}: "
                  f"{len(kontribusi)} file menyumbang baris — pastikan "
                  f"bukan duplikat export (baris dijumlah 2× bila file "
                  f"kembar)")
    return combined


def promote_headers(df: pd.DataFrame) -> pd.DataFrame:
    """
    Mirror M: Table.PromoteHeaders(table, [PromoteAllScalars = true]).
    Row 0 jadi header, drop row 0.
    Jika ada duplicate column names after promote, keep first.
    """
    if df.empty:
        return df

    # Extract row 0 as headers
    headers = df.iloc[0].tolist()

    # Drop row 0
    result = df.iloc[1:].reset_index(drop=True)

    # Set new headers
    # Handle None/NaN in headers → replace with "ColumnN"
    new_cols = []
    seen = {}
    for i, h in enumerate(headers):
        if h is None or (isinstance(h, float) and pd.isna(h)):
            h = f"Column{i+1}"
        h = str(h)
        # Handle duplicates: append _1, _2, etc.
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 0
        new_cols.append(h)

    result.columns = new_cols
    return result


def table_skip(df: pd.DataFrame, n: int) -> pd.DataFrame:
    """Mirror M: Table.Skip(table, n). Skip first n rows."""
    if df.empty or n <= 0:
        return df
    return df.iloc[n:].reset_index(drop=True)


# =============================================================================
# v3.9.33 — PRE-FLIGHT FOLDER DATA (kontrak — dipanggil app.py agent 49-c)
# =============================================================================

def preflight_folders(cfg) -> tuple[list[tuple[str, str]], list[tuple[str, int]]]:
    """v3.9.33 — pre-flight validasi folder data SEBELUM pipeline jalan.

    Return (missing, counts):
    - missing: [(key, full_path)] utk folder Folders[key] yang path-nya
      TIDAK ada (RootPath sendiri invalid → label 'RootPath').
    - counts: [(key, jumlah file .xls*)] per key Folders (folder hilang → 0).

    Audit 48-f HIGH: folder data hilang / RootPath salah dulu membuat
    pipeline 'SELESAI' hijau + tabel kosong SENYAP. app.py memanggil ini
    sebelum run_pipeline utk menampilkan peringatan eksplisit ke user
    (folder mana yang hilang + berapa file yang akan terbaca).
    """
    root = str(getattr(cfg, "RootPath", "") or "")
    missing: list = []
    counts: list = []
    if not os.path.isdir(root):
        # RootPath sendiri invalid → semua subfolder ikut hilang; catat
        # sebagai entri khusus 'RootPath' di depan.
        missing.append(("RootPath", root))
    folders = getattr(cfg, "Folders", None) or {}
    for key in sorted(folders):
        try:
            full = cfg.folder(key)
        except Exception:  # noqa: BLE001 — cfg duck-typed/mock: fallback join
            full = os.path.join(root, str(folders[key]))
        # hitung file utk SEMUA key (folder hilang → 0, kontrak)
        counts.append((key, len(_list_excel_files(full))))
        if not os.path.isdir(full):
            missing.append((key, full))
    return missing, counts


# =============================================================================
# COLUMN OPERATIONS (mirror M Table functions)
# =============================================================================

def table_select_columns(df: pd.DataFrame, columns: List[str]) -> pd.DataFrame:
    """Mirror M: Table.SelectColumns(table, columns, MissingField.Ignore)."""
    available = [c for c in columns if c in df.columns]
    return df[available].copy()


def table_rename_columns(df: pd.DataFrame, rename_map: dict) -> pd.DataFrame:
    """Mirror M: Table.RenameColumns(table, {{old, new}, ...}, MissingField.Ignore)."""
    valid = {old: new for old, new in rename_map.items() if old in df.columns}
    return df.rename(columns=valid)


def safe_to_numeric(series: pd.Series) -> pd.Series:
    """Convert series ke numeric, null untuk yang gagal. Mirror Table.TransformColumnTypes ke number."""
    return pd.to_numeric(series, errors="coerce")


def safe_to_str(series: pd.Series) -> pd.Series:
    """Convert series ke string, null tetap null. Mirror Table.TransformColumnTypes ke text.
    OPTIMIZED: vectorized, no .apply()."""
    # Use vectorized approach: where notna, convert to str; where na, keep None
    mask = series.notna()
    result = pd.Series(index=series.index, dtype=object)
    result[mask] = series[mask].astype(str)
    result[~mask] = None
    return result


def to_datetime_robust(series: pd.Series) -> pd.Series:
    """
    v3.9.28 — pd.to_datetime yang TAHAN SERIAL EXCEL numerik.

    Sumber masalah: file .xlsb dibaca lewat engine pyxlsb menghasilkan sel
    tanggal sebagai SERIAL MENTAH (mis. 46290 = 2026-09-25) — pyxlsb tidak
    mengonversi sel tanggal ke datetime. pd.to_datetime(46290) membacanya
    sebagai NANODETIK-EPOCH (1970-01-01 + 46µs) → SEMUA nilai jatuh ke
    hari 1 → In_H = 0 untuk H > 1 → 'Stok Seharusnya' salah & flag
    'Hitung Manual' palsu (reproduksi: sheet '4  42 05' workbook user —
    4.591 baris serial → 100% hari 1 pra-fix).

    Aturan: nilai NUMERIC dalam rentang serial Excel wajar (20000..80000
    ≈ th 1954..2119 — tak mungkin tumpang tindih dgn nomor hari 1..31)
    dikonversi via epoch Excel 1899-12-30 (pecahan waktu ikut terbawa);
    nilai lain (datetime/Timestamp/string tanggal) diparse
    pd.to_datetime seperti biasa. Kolom datetime64 apa adanya.

    v3.9.42 — AUDIT FIX (61-a, F2): string tanggal kini dayfirst=True.
    Konteks Indonesia = dd/mm/yyyy — tanpa dayfirst, '01/10/2026'
    diparse 10 Jan (padahal maksudnya 1 Okt → H/hari salah) dan
    '13/10/2026' jadi NaT (bulan '13' tak ada → tanggal HILANG senyap).
    Hanya jalur STRING yang kena; jalur serial-Excel numerik di bawah
    tetap apa adanya (numerik bukan parse string).
    format='mixed' WAJIB: inferensi satu-format pandas dgn dayfirst
    membalik kolom ISO ('2026-10-01' ditebak %Y-%d-%m → 10 Jan —
    korupsi senyap); 'mixed' mem-parse per-elemen → dd/mm dan ISO
    hidup berdampingan benar. Perf aman utk data tanggal (nilai unik
    sedikit — per bulan ±31 — cache format pandas efektif).
    """
    s = series if isinstance(series, pd.Series) else pd.Series(series)
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    num = pd.to_numeric(s, errors="coerce")
    serial = num.notna() & num.between(20000.0, 80000.0)
    if not bool(serial.any()):
        # v3.9.42 — dayfirst=True + format='mixed' (lihat docstring):
        # seluruh nilai di sini string/datetime → parse dd/mm/yyyy tanpa
        # merusak format ISO.
        return pd.to_datetime(s, errors="coerce", dayfirst=True,
                              format="mixed")
    # serial Excel (float/int; hasil pyxlsb .xlsb / export numerik) → Timestamp
    ser_ts = pd.Timestamp("1899-12-30") + pd.to_timedelta(
        num.where(serial), unit="D")
    # nilai non-serial → jalur parse biasa (string/datetime objek)
    # v3.9.42 — dayfirst=True + format='mixed' (mirror cabang no-serial
    # di atas): string dd/mm/yyyy di baris non-serial juga harus parse
    # 1 Okt, bukan 10 Jan.
    rest = pd.to_datetime(s.where(~serial), errors="coerce", dayfirst=True,
                          format="mixed")
    return ser_ts.where(serial, rest)


def resto_from_cabang_po(val) -> str:
    """v3.9.30 — Cabang PO Rincian 32.43 → kode resto: bagian setelah
    titik TERAKHIR ("1370.MLGPAR" → "MLGPAR"; tanpa titik → nilai apa
    adanya). Satu logika bersama pipeline (fn_audit PO Menunggu) + UI
    (rincian_po_view._resto_name delegasi ke sini).

    NA-safety (mirror _doc_code_str — dulu `str(val or "")` mengevaluasi
    bool(pd.NA) → TypeError) + float integral di-stringify TANPA ".0"
    (1370.0 → "1370"; dulu "1370.0" → rsplit titik terakhir → "0")."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val or "").strip()
    if not s or s.lower() in ("none", "nan"):
        return ""
    # float non-integral (mis. 1370.5): perilaku lama dipertahankan —
    # rsplit titik terakhir ("1370.5" → "5")
    return s.rsplit(".", 1)[-1].strip() if "." in s else s


def date_to_day(series: pd.Series) -> pd.Series:
    """
    Mirror M: Table.TransformColumns(t, {{"Tanggal", Date.Day, Int64.Type}}).
    Ambil day-of-month dari datetime/date column. Return Int64.
    OPTIMIZED: vectorized via .dt accessor.
    v3.9.28: dibangun di atas to_datetime_robust — kolom tanggal numerik
    (serial Excel, mis. dari file .xlsb via pyxlsb) dikonversi benar,
    bukan dibaca sbg nanodetik-epoch → hari 1 (lihat docstring helper).
    """
    dt = to_datetime_robust(series)
    return dt.dt.day.astype("Int64")
